"""Builds the synthetic workbooks in tests/fixtures. Development only; the tests never run it.

    python -m venv /tmp/fx && /tmp/fx/bin/pip install openpyxl==3.1.5
    /tmp/fx/bin/python tests/build_fixtures.py

The workbooks copy the *shape* of EFRAG's lists (sheet names, a title row above
the header, header labels, merged cells, numeric paragraph cells) so the parser
is tested on realistic structure. Every name and note is invented placeholder
text: no EFRAG content is in these files.
"""
from pathlib import Path

from openpyxl import Workbook

OUT = Path(__file__).resolve().parent / "fixtures"

IG3_HEADER = ["ID", "ESRS", "DR", "Paragraph", "Related AR", "Name", "Data Type", "Conditional or alternative DP",
              "May \n[V]", "Appendix B - ESRS 2 \n(SFDR + PILLAR 3 + Benchmark + CL)",
              "Appendix C - ESRS 1\nPhase-in for undertakings with fewer than 750 employees",
              "Appendix C - ESRS 1\nPhase-in for all undertakings"]
IG3_ESRS2_HEADER = IG3_HEADER[:10] + [
    "Datapoints to be disclosed in case of phased-in [Appendix C - ESRS 1]\nUndertakings with fewer than 750 employees",
    "Appendix C - ESRS 1\n[phased-in datapoints]"]
REVISED_HEADER = ["ID", "Revised ESRS", "DR", "Paragraph", "Related Guidance", "Name",
                  'Disaggregations ("shall" * or "may")', "Data Type", "Conditionality or Alternative presentation",
                  "Conditionality defined in a different paragraph",
                  "Appendix A - ESRS 2 \n(SFDR + PILLAR 3 + BENCHMARK + CL)",
                  "Phase-ins for 'Wave-one' undertakings exceeding the thresholds",
                  "Phase-ins for 'Wave-one' undertakings NOT exceeding the thresholds",
                  "Phase-ins for 'Other undertakings'"]
MAPPING_HEADER = ["Mapping with 2024 IG 3 (2023 ESRS ID)", "2023 ESRS", "2023 ESRS DR", "2023 ESRS Paragraph",
                  "2023 ESRS related AR"]

IG3 = {
    "ESRS 2": [
        ["BP-1_01", "ESRS 2", "BP-1", "5 a", "", "Placeholder datapoint 10 about the basis for preparation",
         "semi-narrative", "", "", "", "", ""],
        ["SBM-3_01", "ESRS 2", "SBM-3", "48 a", "AR 16", "Placeholder datapoint 11 on material impacts", "narrative",
         "", "", "", "According to the phase-in appendix (placeholder)", ""],
        ["IRO-1_01", "ESRS 2", "IRO-1", 53, "", "Placeholder datapoint 12 about the assessment process", "narrative",
         "Conditional", "V", "", "", "1 year"],
    ],
    "ESRS 2 MDR": [
        ["MDR-P_01", "ESRS 2", "MDR-P", "65 a", "AR 21", "Placeholder datapoint 13 about policy contents",
         "narrative", "", " ", " ", ""],
        ["MDR-P_02", "ESRS 2", "MDR-P", "65 b", "", "Placeholder datapoint 14 about policy scope", "narrative",
         "", "", "", ""],
        None,  # blank row, then a section note alone in the ID column
        ["Placeholder section note: the datapoints below apply only in a stated case"],
        ["MDR-P_07", "ESRS 2", "MDR-P", 62, "", "Placeholder datapoint 15 about reasons for having no policy",
         "narrative", "Conditional", "", "", ""],
    ],
    "ESRS E1": [
        ["E1-1_01", "E1", "E1-1", 14, "", "Placeholder datapoint 1 about a transition plan", "narrative",
         "", "", "", "", ""],
        ["E1-1_02", "E1", "E1-1", "16 a", "AR 1", "Placeholder datapoint 2 about locked-in items",
         "semi-narrative", "Conditional", "V", "CL", "", ""],
        ["E1.MDR-P_01-02", "ESRS 2", "", 62, "", "Placeholder reference row to the minimum disclosure requirements",
         "MDR-P", "", "", "", "", ""],
        ["E1-6_01", "E1", " E1-6 ", "44 a", "", "Placeholder datapoint 4 on gross scope 1 emissions", "ghgEmissions",
         "", "", "SFDR+PILLAR 3+ BENCHMARK", "", ""],
        ["E1-6_02", "E1", "E1-6", "44 c", "AR 46", "Placeholder datapoint 5 on gross scope 3 emissions",
         "ghgEmissions", "", "", "SFDR", "1 year", "3 years"],
        ["E1-6_03", "E1", "E1-6", "AR 46 d", "", "Placeholder datapoint 6: Scope3 category breakdown (invented)",
         "Table/ghgEmissions", "Alternative", "V", "", "1 year", ""],
        ["E1-7_01", "E1", "E1-7", 56, "", "Placeholder datapoint 9 that is dropped later", "narrative",
         "", "", "", "", ""],
        ["E1-9_01", "E1", "E1-9", "64 a", "AR 70", "Placeholder datapoint 7 about anticipated effects",
         "monetary", "", "", "BENCHMARK", "", "3 years"],
        ["E1-9_02", "E1", "E1-9", 66, "", "Placeholder datapoint 8 — CO₂ résumé naïve "
         "unicode test", "percent", "", "V", "", "", ""],
    ],
    "ESRS S1": [
        ["S1-6_01", "S1", "S1-6", "50 a", "", "Placeholder datapoint 30 about headcount", "Integer", "", "", "", "",
         "1 year"],
        ["S1-6_02", "S1", "S1-6", "50 b", "", "Placeholder datapoint 31 about water in dormitories (invented)",
         "narrative", "", "V", "SFDR/BENCH", "", ""],
    ],
}

# (row values, mapping columns) for the revised layout.
REVISED = {
    "ESRS 2": [
        (["ESRS26_BP-1_01", "ESRS 2", "BP-1", "4 a", "", "Placeholder datapoint 10 about the basis for preparation",
          "", "semi-narrative", "", "", "", "", "", ""], ["BP-1_01", "ESRS 2", "BP-1", "5 a", ""]),
        (["ESRS26_SBM-3_01", "ESRS 2", "SBM-3", "33 a", "AR 20", "Placeholder datapoint 11 on material impacts, "
          "reworded", "Impact, risk or opportunity", "narrative", "Conditional", "28, 29", "", "", "", ""],
         ["SBM-3_01", "ESRS 2", "SBM-3", "48 a", "AR 16"]),
    ],
    "ESRS 2 GDR": [
        (["ESRS26_GDR-P_01", "ESRS 2", "GDR-P", "42 a", "", "Placeholder datapoint 16 about policy contents",
          "Policy(ies)", "narrative", "", "", "", "", "", ""], ["MDR-P_01, MDR-P_02", "ESRS 2", "MDR-P", "65", ""]),
    ],
    "ESRS E1": [
        (["ESRS26_E1-1_01", "ESRS E1", "E1-1", "12 a", "AR 1", "Placeholder datapoint 1 about a transition plan", "",
          "narrative", "", "", "CL", "", "prior to FY 2027", ""], ["E1-1_01", "E1", "E1-1", "14", ""]),
        (["ESRS26_E1-8_01", "ESRS E1", "E1-8", "24 a", "", "Placeholder datapoint 4 on gross scope 1 emissions",
          "Gas type*", "ghgemissions", "", "", "SFDR + PILLAR 3 + BENCHMARK", "", "prior to FY 2027", ""],
         ["E1-6_01", "E1", "E1-6", "44 a", ""]),
        (["ESRS26_E1-8_02", "ESRS E1", "E1-8", "24 c", "", "Placeholder datapoint 5 on gross scope 3 emissions and "
          "categories", "Scope 3 category*", "ghgemissions", "Conditional", "17", "SFDR", "prior to FY 2030",
          "prior to FY 2030", "first four FY of reporting"], ["E1-6_02, E1-6_03", "E1", "E1-6", "44 c", "AR 46"]),
        (["E1-9_01", "ESRS E1", "E1-11", "39 a", "", "Placeholder datapoint 7 about anticipated effects, reworded",
          "", "monetary", "", "", "BENCHMARK", "prior to FY 2028", "prior to FY 2028", "first two FY of reporting"],
         ["", "", "", "", ""]),
        (["ESRS26_E1-11_02", "ESRS E1", "E1-11", "40", "", "Placeholder datapoint 8 — CO₂ résumé "
          "naïve unicode test", "", "percent", "", "", "", "", "prior to FY 2027", ""], ["", "", "", "", ""]),
        (["ESRS26_E1-2_01", "ESRS E1", "E1-2", "15", "", "Placeholder datapoint 20 that is new in this version", "",
          "narrative", "", "", "", "", "prior to FY 2027", ""], ["", "", "", "", ""]),
        (["ESRS26_E1.GDR-P", "ESRS E1", "E1-4", "20", "", "Placeholder reference to the general disclosure "
          "requirements for policies", "", "GDR-P", "", "", "", "", "prior to FY 2027", ""],
         ["MDR-P_07, E1-4_99", "ESRS 2", "MDR-P", "62", ""]),
    ],
    "ESRS S1": [
        (["ESRS26_S1-5_01", "ESRS S1", "S1-5", "30 a", "", "Placeholder datapoint 30 about headcount", "Gender*",
          "integer", "Alternative presentation", "", "SFDR", "prior to FY 2027", "prior to FY 2027",
          "first FY of reporting"], ["S1-6_01", "S1", "S1-6", "50 a", ""]),
    ],
}


def title_row(ws, text, width):
    ws.append([text])
    ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=width)


def ig3_style():
    wb = Workbook()
    ws = wb.active
    ws.title = "Index"
    ws["B3"] = "Synthetic test workbook in the IG 3 layout (invented placeholder text)"
    ws["B5"] = "Placeholder legend line"
    for name, rows in IG3.items():
        ws = wb.create_sheet(name)
        header = IG3_ESRS2_HEADER if name == "ESRS 2" else IG3_HEADER[:11] if name.endswith("MDR") else IG3_HEADER
        title_row(ws, "INSTRUCTIONS\nPlaceholder instructions for this synthetic sheet", len(header))
        ws.append(header)
        for r in rows:
            if r is None:
                ws.append([])
            elif len(r) == 1:
                ws.append(r)
                ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=len(header))
            else:
                ws.append(r)
    wb.save(OUT / "ig3_style.xlsx")


def revised_style(mapping: bool):
    wb = Workbook()
    ws = wb.active
    ws.title = "Disclaimer"
    title_row(ws, "Placeholder disclaimer for a synthetic workbook", 10)
    ws = wb.create_sheet("Index")
    ws["B2"] = "Version: 1 January 2030"
    ws["B5"] = "Synthetic test workbook in the revised-ESRS layout (invented placeholder text)"
    for name, rows in REVISED.items():
        ws = wb.create_sheet(name)
        header = (MAPPING_HEADER if mapping else []) + REVISED_HEADER
        title_row(ws, "Placeholder disclaimer row", len(header))
        ws.append(header)
        for values, mapped in rows:
            ws.append((mapped if mapping else []) + values)
    ws = wb.create_sheet("Statistics")
    ws["C3"] = "Placeholder statistics"
    ws["C5"], ws["D5"] = "ESRS", "Total"
    wb.save(OUT / ("revised_style_mapping.xlsx" if mapping else "revised_style_clean.xlsx"))


def header_variants():
    wb = Workbook()
    ws = wb.active
    ws.title = "Datapoints"
    title_row(ws, "Placeholder list with its header on row 5 and its own column names", 8)
    ws.append([])
    ws.append(["Placeholder note row"])
    ws.append([])
    ws.append(["Datapoint name", "Comment", "Datapoint ID", "Disclosure Requirement", "Paragraph reference",
               "Datatype", "Voluntary", "Standard"])
    ws.append(["Placeholder datapoint A1 — Ünïcödé naïve café", "free comment",
               "G1-1_01", "G1-1", 7, "narrative", "", "G1"])
    ws.append(["Placeholder datapoint A2 about suppliers", "", "G1-1_02", None, 8, "semi-narrative", "x", "G1"])
    ws.append(["Placeholder datapoint A3 about payment terms", "", "G1-1_03", None, "8 b", "monetary", "", "ESRS G1"])
    ws.merge_cells("D6:D8")  # one DR code for three datapoint rows
    ws.append(["Placeholder row without an ID", "", None, "G1-2", 9, "narrative", "", "G1"])
    ws.append(["Placeholder datapoint A5, same ID as A1", "", "G1-1_01", "G1-1", 10, "narrative", "", "G1"])
    ws.append(["Placeholder datapoint A6", "", "not a valid id!", "G1-2", 11, "narrative", "", "G1"])
    ws.append([])
    ws.append(["Datapoint name", "Comment", "Datapoint ID", "Disclosure Requirement"])
    ws.append(["Placeholder datapoint A7 after a repeated header", "", "G1-3_01", "G1-3", "12", "percent", "V", "G1"])
    hidden = wb.create_sheet("Old copy")
    hidden.append(["ID", "Name"])
    hidden.append(["G1-9_01", "Placeholder datapoint in a hidden sheet"])
    hidden.sheet_state = "hidden"
    notes = wb.create_sheet("Notes")
    notes["A1"] = "Placeholder notes sheet without a table"
    wb.save(OUT / "header_variants.xlsx")


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    ig3_style()
    revised_style(False)
    revised_style(True)
    header_variants()
    print("written:", ", ".join(sorted(p.name for p in OUT.glob("*.xlsx"))))
