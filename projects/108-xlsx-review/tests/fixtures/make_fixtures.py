"""Regenerate the openpyxl-written fixtures. Not part of the test run and not imported by the tool.

    python3 -m venv /tmp/venv && /tmp/venv/bin/pip install openpyxl==3.1.5
    /tmp/venv/bin/python tests/fixtures/make_fixtures.py

budget_before.xlsx             a small six-year budget model: inputs, model, summary, defined names
budget_after.xlsx              the same model after an agent's edit: a row inserted (references adjusted, as Excel
                               does), one formula replaced by a number, one formula typed differently from its row,
                               a sheet renamed, a defined name repointed, an external link added, and a bonus pool
                               that makes net income depend on itself across two sheets
budget_after_insert_rows.xlsx  the model after openpyxl's insert_rows(6), which does not adjust formulas
                               ("Openpyxl does not manage dependencies, such as formulae ... when rows or columns are
                               inserted or deleted", https://openpyxl.readthedocs.io/en/stable/editing_worksheets.html)

The external link is added after saving, by editing the zip, because openpyxl has no public API to create one.
"""
import io
import sys
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName

HERE = Path(__file__).resolve().parent
YEARS = 6
COLS = "BCDEFG"


def build_before() -> Workbook:
    wb = Workbook()
    inputs = wb.active
    inputs.title = "Inputs"
    rows = [("Assumption", "Value"), ("Revenue 2025 (EUR k)", 1200), ("Revenue growth", 0.08),
            ("Cost of sales ratio", 0.42), ("Opex growth", 0.03), ("Opex 2025 (EUR k)", 310), ("Tax rate", 0.25)]
    for r, (label, value) in enumerate(rows, 1):
        inputs.cell(r, 1, label)
        inputs.cell(r, 2, value)
    for name, ref in (("StartRevenue", "$B$2"), ("Growth", "$B$3"), ("CostRatio", "$B$4"), ("OpexGrowth", "$B$5"),
                      ("StartOpex", "$B$6"), ("TaxRate", "$B$7")):
        wb.defined_names[name] = DefinedName(name, attr_text=f"Inputs!{ref}")

    model = wb.create_sheet("Model")
    model["A1"] = "EUR k"
    model["B1"] = 2025
    labels = ["Revenue", "Cost of sales", "Gross profit", "Opex", "EBITDA", "Tax", "Net income", None, "Gross margin"]
    for r, label in enumerate(labels, 2):
        if label:
            model.cell(r, 1, label)
    for i, col in enumerate(COLS):
        prev = COLS[i - 1] if i else None
        if prev:
            model[f"{col}1"] = f"={prev}1+1"
            model[f"{col}2"] = f"={prev}2*(1+Growth)"
            model[f"{col}5"] = f"={prev}5*(1+OpexGrowth)"
        else:
            model[f"{col}2"] = "=StartRevenue"
            model[f"{col}5"] = "=StartOpex"
        model[f"{col}3"] = f"={col}2*CostRatio"
        model[f"{col}4"] = f"={col}2-{col}3"
        model[f"{col}6"] = f"={col}4-{col}5"
        model[f"{col}7"] = f"=MAX(0,{col}6*TaxRate)"
        model[f"{col}8"] = f"={col}6-{col}7"
        model[f"{col}10"] = f"=IF({col}2=0,0,{col}4/{col}2)"

    summary = wb.create_sheet("Summary")
    for r, (label, formula) in enumerate([("Summary (EUR k)", None), ("Revenue 2025-2030", "=SUM(Model!B2:G2)"),
                                          ("EBITDA 2025-2030", "=SUM(Model!B6:G6)"), ("Net income 2030", "=Model!G8"),
                                          ("Average gross margin", "=AVERAGE(Model!B10:G10)")], 1):
        summary.cell(r, 1, label)
        if formula:
            summary.cell(r, 2, formula)
    return wb


def build_after() -> Workbook:
    """The agent's version, with references adjusted for the inserted row as Excel would adjust them."""
    wb = build_before()
    model = wb["Model"]
    # Insert "Marketing" as row 6 and rewrite every formula from row 6 down as Excel would after an insert.
    for r in range(11, 5, -1):
        for c in range(1, 8):
            model.cell(r + 1, c).value = model.cell(r, c).value
            model.cell(r, c).value = None
    model["A6"] = "Marketing"
    for i, col in enumerate(COLS):
        model[f"{col}6"] = 40 + 5 * i
        model[f"{col}7"] = f"={col}4-{col}5-{col}6"       # EBITDA now subtracts marketing
        model[f"{col}8"] = f"=MAX(0,{col}7*TaxRate)"
        model[f"{col}9"] = f"={col}7-{col}8"
        model[f"{col}11"] = f"=IF({col}2=0,0,{col}4/{col}2)"
    model["F2"] = 1714                                      # a pasted number where the growth formula was
    model["E3"] = "=E2*0.42"                                # the cost ratio typed in, not taken from Inputs
    inputs = wb["Inputs"]
    inputs["A8"] = "Tax rate from 2027"
    inputs["B8"] = 0.28
    wb.defined_names["TaxRate"] = DefinedName("TaxRate", attr_text="Inputs!$B$8")
    summary = wb["Summary"]
    summary.title = "Dashboard"
    summary["B3"] = "=SUM(Model!B7:G7)"
    summary["B4"] = "=Model!G9"
    summary["B5"] = "=AVERAGE(Model!B11:G11)"
    summary["A6"] = "Bonus pool (10% of 2030 net income)"
    summary["B6"] = "=Model!G9*0.1"
    model["G8"] = "=MAX(0,(G7-Dashboard!B6)*TaxRate)"     # tax after the bonus: the loop
    summary["A7"] = "EUR/USD (from the treasury file)"
    summary["B7"] = "=[1]Rates!$B$2"
    return wb


def build_insert_rows() -> Workbook:
    wb = build_before()
    model = wb["Model"]
    model.insert_rows(6)
    model["A6"] = "Marketing"
    for i, col in enumerate(COLS):
        model[f"{col}6"] = 40 + 5 * i
    return wb


def save(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def add_external_link(data: bytes, target: str) -> bytes:
    """Add xl/externalLinks/externalLink1.xml pointing at target, as ECMA-376 §18.14 and §18.2.8 describe."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist():
            body = src.read(info.filename)
            if info.filename == "xl/workbook.xml":
                text = body.decode("utf-8")
                text = text.replace("</sheets>", '</sheets><externalReferences><externalReference r:id="rIdLink1"/>'
                                                  "</externalReferences>", 1)
                body = text.encode("utf-8")
            elif info.filename == "xl/_rels/workbook.xml.rels":
                text = body.decode("utf-8")
                text = text.replace("</Relationships>", f'<Relationship Id="rIdLink1" Type="{rns}/externalLink" '
                                                        'Target="externalLinks/externalLink1.xml"/></Relationships>')
                body = text.encode("utf-8")
            elif info.filename == "[Content_Types].xml":
                text = body.decode("utf-8")
                text = text.replace("</Types>", '<Override PartName="/xl/externalLinks/externalLink1.xml" ContentType='
                                                '"application/vnd.openxmlformats-officedocument.spreadsheetml.externalLink+xml"/>'
                                                "</Types>")
                body = text.encode("utf-8")
            z.writestr(info, body)
        z.writestr("xl/externalLinks/externalLink1.xml",
                   f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<externalLink xmlns="{ns}" xmlns:r="{rns}">'
                   '<externalBook r:id="rId1"><sheetNames><sheetName val="Rates"/></sheetNames>'
                   '<sheetDataSet><sheetData sheetId="0"><row r="2"><cell r="B2"><v>1.0842</v></cell></row></sheetData>'
                   "</sheetDataSet></externalBook></externalLink>")
        z.writestr("xl/externalLinks/_rels/externalLink1.xml.rels",
                   '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns='
                   '"http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" '
                   f'Type="{rns}/externalLinkPath" Target="{target}" TargetMode="External"/></Relationships>')
    return out.getvalue()


def main() -> int:
    (HERE / "budget_before.xlsx").write_bytes(save(build_before()))
    (HERE / "budget_after.xlsx").write_bytes(
        add_external_link(save(build_after()), "file:///C:/Users/analyst/Downloads/fx_rates.xlsx"))
    (HERE / "budget_after_insert_rows.xlsx").write_bytes(save(build_insert_rows()))
    for name in ("budget_before.xlsx", "budget_after.xlsx", "budget_after_insert_rows.xlsx"):
        load_workbook(HERE / name)  # openpyxl can read back what it wrote
        print(name, (HERE / name).stat().st_size, "bytes")
    return 0


if __name__ == "__main__":
    sys.exit(main())
